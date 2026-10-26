"""Validation du rapport Excel produit par le notebook et synthèse des indicateurs.

Ce module ne recalcule AUCUN statut métier : il lit les objets produits par le moteur
(``MATCH_RUN_SUMMARY``, ``AUDIT``, ``STRUCTURE_RUN_SUMMARY``) et le fichier
``global/customer_controls.csv`` écrit par ``export_reports`` pendant le même run.
Le classeur téléchargé est le fichier exact écrit par le notebook, jamais une reconstruction.
"""
from __future__ import annotations

import csv
import hashlib
import os
from dataclasses import dataclass
from pathlib import Path

from config import AppConfig
from exceptions import ReportGenerationError

REQUIRED_SHEETS = ("SUMMARY", "CUSTOMER_CONTROLS", "DOCUMENT_CONTROLS", "DATABASE_MATCHING",
                   "FATCA_DETAILS", "SIGNATURE_DETAILS")
MTIME_TOLERANCE_S = 2.0
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass
class ReportInfo:
    path: Path
    size_bytes: int
    mtime_ns: int
    sha256: str
    sheets: list[str]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(4 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_report(path: Path, not_before_epoch: float) -> ReportInfo:
    """Le rapport doit exister, être non vide, écrit pendant ce run et lisible par openpyxl."""
    if not path.exists():
        raise ReportGenerationError(f"Rapport absent : {path}")
    info = path.stat()
    if info.st_size <= 0:
        raise ReportGenerationError(f"Rapport vide : {path}")
    if info.st_mtime < not_before_epoch - MTIME_TOLERANCE_S:
        raise ReportGenerationError(
            f"Rapport antérieur au run (mtime {info.st_mtime:.0f} < début {not_before_epoch:.0f}) : ancien fichier non régénéré")
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True)
        sheets = list(workbook.sheetnames)
        workbook.close()
    except Exception as exc:  # fichier tronqué ou non-xlsx
        raise ReportGenerationError(f"Rapport illisible : {type(exc).__name__}: {exc}") from exc
    missing = [s for s in REQUIRED_SHEETS if s not in sheets]
    if missing:
        raise ReportGenerationError("Feuilles manquantes dans le rapport : " + ", ".join(missing))
    return ReportInfo(path=path, size_bytes=info.st_size, mtime_ns=info.st_mtime_ns, sha256=_sha256(path), sheets=sheets)


def read_report_for_download(path: Path, expected_sha256: str | None, expected_mtime_ns: int | None) -> bytes | None:
    """Octets du rapport, seulement s'il est identique à celui validé à la fin du run."""
    if not expected_sha256 or not path.is_file():
        return None
    try:
        if expected_mtime_ns is not None and path.stat().st_mtime_ns != expected_mtime_ns:
            return None
        data = path.read_bytes()
    except OSError:
        return None
    return data if hashlib.sha256(data).hexdigest() == expected_sha256 else None


def _read_customer_controls(cfg: AppConfig, not_before_epoch: float) -> list[dict]:
    path = cfg.output_root / "global" / "customer_controls.csv"
    try:
        if path.stat().st_mtime < not_before_epoch - MTIME_TOLERANCE_S:
            return []
        with path.open(encoding="utf-8", newline="") as stream:
            return list(csv.DictReader(stream))
    except (OSError, csv.Error):
        return []


def build_summary(core, cfg: AppConfig, not_before_epoch: float) -> dict:
    """Indicateurs de gestion issus des objets et fichiers produits par le notebook pendant ce run."""
    match = getattr(core, "MATCH_RUN_SUMMARY", {}) or {}
    counts = match.get("counts", {}) or {}
    audit = getattr(core, "AUDIT", {}) or {}
    physical = audit.get("PHYSICAL_DOCUMENTS", {}) or {}
    raw = audit.get("RAW_OCR", {}) or {}
    structure = getattr(core, "STRUCTURE_RUN_SUMMARY", {}) or {}
    raw_run = (getattr(core, "RAW_RUN_SUMMARY", {}) or {}).get("summary", {}) or {}
    controls = _read_customer_controls(cfg, not_before_epoch)
    manual = sum(str(row.get("MANUAL_RECHECK_REQUIRED", "")).strip().lower() == "true" for row in controls)
    errors = list(getattr(core, "ERRORS", []) or [])
    selections = list(getattr(core, "SELECTED_DOCUMENTS", []) or [])
    selected_ids = set(getattr(core, "CUSTOMERS_SELECTED", []) or [])
    scoped = [s for s in selections if s["customer_id"] in selected_ids]
    anomalies = []
    if raw_run.get("customers_failed") or raw_run.get("pages_failed"):
        anomalies.append(f"RAW OCR : {raw_run.get('pages_failed', 0)} page(s) en échec enregistrée(s)")
    if structure.get("status") == "COMPLETED_WITH_ERRORS":
        anomalies.append(f"Extraction structurée : {len(structure.get('customers_failed', []))} client(s) avec erreurs")
    if match.get("customers_failed"):
        anomalies.append(f"Rapprochement : {len(match['customers_failed'])} client(s) en erreur de traitement")
    if errors:
        anomalies.append(f"{len(errors)} erreur(s) technique(s) journalisée(s) par le pipeline")
    return {
        "customers_discovered": len(getattr(core, "ALL_CUSTOMERS", []) or []),
        "customers_selected": len(selected_ids),
        "customers_processed": len(match.get("customers_completed", []) or []),
        "documents_processed": int(physical.get("document_jsons", physical.get("selected", 0)) or 0),
        "documents_missing": sum(not s.get("document_found") for s in scoped),
        "duplicates_skipped": int(physical.get("skipped_same_family_same_size", 0) or 0),
        "raw_pages_processed": int(raw.get("pages_completed", 0) or 0),
        "conform_count": int(counts.get("CONFORME_A_LA_BASE", 0)),
        "non_conform_count": int(counts.get("NON_CONFORME_A_LA_BASE", 0)),
        "recheck_count": int(counts.get("A_RECONTROLER", 0)),
        "reference_not_found_count": int(counts.get("REFERENCE_NOT_FOUND", 0)),
        "processing_error_count": int(counts.get("PROCESSING_ERROR", 0)),
        "manual_review_count": manual,
        "qwen_calls": int(getattr(getattr(core, "BUDGET", None), "used", 0) or 0),
        "anomalies": anomalies,
        "technical": {
            "raw_ocr": raw, "structure_status": structure.get("status"),
            "matching_counts": counts, "name_similarity": audit.get("NAME_SIMILARITY", {}),
            "pipeline_errors": len(errors), "engine_session_id": getattr(core, "SESSION_ID", None),
            "performance": audit.get("PERFORMANCE", {}), "outputs": audit.get("OUTPUTS", {}),
        },
    }


def describe_report_file(path: Path) -> dict:
    try:
        info = path.stat()
    except OSError:
        return {"exists": False}
    return {"exists": True, "size_bytes": info.st_size, "mtime_epoch": info.st_mtime,
            "readable": os.access(path, os.R_OK)}
