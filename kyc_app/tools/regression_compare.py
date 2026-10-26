"""Comparaison de non-régression : sorties du NOTEBOOK vs sorties de l'APPLICATION.

    python tools/regression_compare.py --reference /mnt/output_qwen38_notebook \
                                       --candidate /mnt/output_qwen38_app_regression \
                                       [--json /mnt/kyc_regression/comparison.json]

Compare, sur la même archive :
  * population : clients découverts, PDF sélectionnés / manquants / doublons ;
  * statuts clients (customer_controls.csv) : familles, matchings, revue, cohérence base ;
  * statuts documents (document_controls.csv) : OCR, type, biométrie, expiration, valeurs, statut base ;
  * similarité FATCA et CARTON SIGNATURE (scores non arrondis, décision, fiabilité, raison) ;
  * synthèse similarité (name_similarity_summary.json) ;
  * classeur Excel : feuilles présentes et nombre de lignes des feuilles métier.
Les colonnes volatiles (chemins absolus, horodatages, identifiants de session, performance)
sont exclues. Code retour 0 = résultats métier équivalents, 1 = écarts, 2 = fichiers manquants.
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

CUSTOMER_COLUMNS = ["REFERENCE_STATUS", "IDENTITY_STATUS", "DOMICILE_STATUS", "CONVENTION_STATUS", "FATCA_STATUS",
                    "SIGNATURE_STATUS", "MATCH_NAME", "MATCH_DATE_OF_BIRTH", "MATCH_EXPIRATION_DATE",
                    "MATCH_ACCOUNT_NUMBER", "MANUAL_RECHECK_REQUIRED", "DATABASE_CONSISTENCY_STATUS", "ANOMALY_CODES"]
DOCUMENT_COLUMNS = ["LOGICAL_DOCUMENT_FAMILY", "PHYSICAL_FILENAME", "FILE_SIZE", "SHA256", "OCR_STATUS", "DOCUMENT_TYPE",
                    "BIOMETRIC_STATUS", "EXPIRATION_STATUS", "NAME_EXTRACTED", "DOB_EXTRACTED", "EXPIRATION_EXTRACTED",
                    "ACCOUNT_EXTRACTED", "DOCUMENT_DATABASE_STATUS", "MANUAL_RECHECK_REQUIRED"]
SIMILARITY_COLUMNS = ["FORM_TYPE", "EXTRACTED_NAME", "REFERENCE_NAME", "EXTRACTED_NAME_NORMALIZED",
                      "REFERENCE_NAME_NORMALIZED", "EXACT_MATCH", "DIRECT_SIMILARITY", "TOKEN_ORDER_SIMILARITY",
                      "FINAL_SIMILARITY_SCORE", "THRESHOLD", "SIMILARITY_MATCH", "NAME_RELIABLE", "NAME_MATCH_STATUS",
                      "NAME_CONFORMITY_REASON", "DOCUMENT_DATABASE_STATUS", "CONFORMITY_REASON",
                      "SIGNATURE_ACCOUNT_MATCH", "SIGNATURE_ACCOUNT_RELIABLE_MATCH"]
NUMERIC_COLUMNS = {"DIRECT_SIMILARITY", "TOKEN_ORDER_SIMILARITY", "FINAL_SIMILARITY_SCORE", "THRESHOLD"}
KEY_SHEETS = ["CUSTOMER_CONTROLS", "DOCUMENT_CONTROLS", "DOCUMENT_INVENTORY", "DUPLICATE_GROUPS", "DATABASE_MATCHING",
              "NON_CONFORMITIES", "MANUAL_REVIEW_QUEUE", "LOW_CONFIDENCE_FIELDS", "HANDWRITTEN_FIELDS",
              "EXTRACTION_COVERAGE", "DOCUMENT_TYPE_COVERAGE", "IDENTITY_DETAILS", "DOMICILE_DETAILS",
              "CONVENTION_DETAILS", "FATCA_DETAILS", "SIGNATURE_DETAILS", "NAME_SIMILARITY_DISTRIBUTION",
              "NAME_SIMILARITY_THRESHOLDS"]
VOLATILE_SHEETS = {"SUMMARY", "PERFORMANCE", "PROCESSING_ERRORS"}


class Comparison:
    def __init__(self) -> None:
        self.checks: list[dict] = []

    def add(self, area: str, key: str, field: str, reference, candidate) -> None:
        equal = reference == candidate
        if not equal and field in NUMERIC_COLUMNS:
            try:
                equal = abs(float(reference) - float(candidate)) <= 1e-9
            except (TypeError, ValueError):
                pass
        self.checks.append({"area": area, "key": key, "field": field, "equal": equal,
                            "reference": reference, "candidate": candidate})

    @property
    def differences(self) -> list[dict]:
        return [c for c in self.checks if not c["equal"]]


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def normalize_anomalies(value: str) -> str:
    try:
        return json.dumps(sorted(json.loads(value)), ensure_ascii=False)
    except (TypeError, ValueError):
        return value


def compare_table(cmp: Comparison, area: str, reference: list[dict], candidate: list[dict], key_columns: list[str],
                  columns: list[str]) -> None:
    def index(rows):
        return {"|".join(str(r.get(k, "")) for k in key_columns): r for r in rows}

    ref, cand = index(reference), index(candidate)
    cmp.add(area, "*", "row_count", len(reference), len(candidate))
    for key in sorted(set(ref) | set(cand)):
        if key not in ref or key not in cand:
            cmp.add(area, key, "row_present", key in ref, key in cand)
            continue
        for column in columns:
            if column not in ref[key] and column not in cand[key]:
                continue
            left, right = ref[key].get(column), cand[key].get(column)
            if column == "ANOMALY_CODES":
                left, right = normalize_anomalies(left), normalize_anomalies(right)
            cmp.add(area, key, column, left, right)


def population_summary(root: Path) -> dict:
    data = json.loads((root / "inventory" / "population.json").read_text(encoding="utf-8"))
    selections = data["selections"]
    return {"customers": sorted(data["customers"]),
            "documents_selected": sum(bool(s["document_found"]) for s in selections),
            "documents_missing": sum(not s["document_found"] for s in selections),
            "duplicates_skipped": sum(int(s["duplicate_count"]) for s in selections),
            "selection_identity": sorted((s["customer_id"], s["document_key"], s["physical_document_id"],
                                          s["pdf_sha256"] or "") for s in selections)}


def sheet_rows(path: Path) -> dict:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True)
    rows = {sheet.title: sheet.max_row for sheet in workbook}
    workbook.close()
    return rows


def run(reference: Path, candidate: Path) -> Comparison:
    cmp = Comparison()
    left_pop, right_pop = population_summary(reference), population_summary(candidate)
    for field in ("customers", "documents_selected", "documents_missing", "duplicates_skipped", "selection_identity"):
        cmp.add("population", "*", field, left_pop[field], right_pop[field])
    compare_table(cmp, "customers", read_csv(reference / "global/customer_controls.csv"),
                  read_csv(candidate / "global/customer_controls.csv"), ["CUSTOMER_ID"], CUSTOMER_COLUMNS)
    compare_table(cmp, "documents", read_csv(reference / "global/document_controls.csv"),
                  read_csv(candidate / "global/document_controls.csv"), ["CUSTOMER_ID", "PHYSICAL_DOCUMENT_ID"],
                  DOCUMENT_COLUMNS)
    for name in ("fatca_name_matching.csv", "signature_name_matching.csv"):
        compare_table(cmp, name.removesuffix(".csv"), read_csv(reference / "global" / name),
                      read_csv(candidate / "global" / name), ["CUSTOMER_ID", "PHYSICAL_DOCUMENT_ID", "FORM_INDEX"],
                      SIMILARITY_COLUMNS)
    left_sim = json.loads((reference / "global/name_similarity_summary.json").read_text(encoding="utf-8"))
    right_sim = json.loads((candidate / "global/name_similarity_summary.json").read_text(encoding="utf-8"))
    cmp.add("name_similarity_summary", "*", "content", left_sim, right_sim)
    report = "reports/kyc_recertification_control.xlsx"
    left_rows, right_rows = sheet_rows(reference / report), sheet_rows(candidate / report)
    cmp.add("excel", "*", "sheets", sorted(left_rows), sorted(right_rows))
    for sheet in KEY_SHEETS:
        cmp.add("excel", sheet, "rows", left_rows.get(sheet), right_rows.get(sheet))
    return cmp


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--reference", type=Path, required=True, help="OUTPUT_DIR de l'exécution du notebook")
    parser.add_argument("--candidate", type=Path, required=True, help="OUTPUT_ROOT de l'exécution de l'application")
    parser.add_argument("--json", type=Path, help="écrire le détail des contrôles en JSON")
    args = parser.parse_args(argv)
    try:
        cmp = run(args.reference, args.candidate)
    except (FileNotFoundError, KeyError, ValueError) as exc:
        print("COMPARAISON IMPOSSIBLE :", exc, file=sys.stderr)
        return 2
    areas: dict[str, list[int]] = {}
    for check in cmp.checks:
        totals = areas.setdefault(check["area"], [0, 0])
        totals[0] += 1
        totals[1] += int(not check["equal"])
    print(f"{'domaine':32} {'contrôles':>10} {'écarts':>8}")
    for area, (count, diffs) in areas.items():
        print(f"{area:32} {count:>10} {diffs:>8}")
    for diff in cmp.differences[:50]:
        print(f"ÉCART {diff['area']} [{diff['key']}] {diff['field']} : notebook={diff['reference']!r} app={diff['candidate']!r}")
    if args.json:
        args.json.write_text(json.dumps({"checks": cmp.checks, "differences": cmp.differences},
                                        ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    verdict = "ÉQUIVALENT" if not cmp.differences else f"{len(cmp.differences)} ÉCART(S)"
    print(f"\nRÉSULTAT : {verdict} ({len(cmp.checks)} contrôles)")
    return 0 if not cmp.differences else 1


if __name__ == "__main__":
    sys.exit(main())
